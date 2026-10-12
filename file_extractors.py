"""
file_extractors.py — PDF, Word, Excel, PPT, images, text se text nikaalo.
LLM/RAG ke liye sab formats ek hi text pipeline me aate hain.
"""

import csv
import html
import io
import json
import re
from typing import Callable, Dict, List

from config import GEMINI_API_KEY, logger

EXTRACTABLE_EXTENSIONS = [
    ".pdf", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".rtf",
    ".docx", ".doc",
    ".xlsx", ".xls", ".xlsm",
    ".pptx",
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".tif",
]


def ext_of(filename: str) -> str:
    name = (filename or "").lower()
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1]


def extract_pdf(data: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for i, page in enumerate(reader.pages):
        parts.append(f"--- Page {i + 1} ---\n{page.extract_text() or ''}")
    return "\n".join(parts)


def extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_csv(data: bytes) -> str:
    text = extract_txt(data)
    reader = csv.reader(io.StringIO(text))
    rows = []
    for row in reader:
        rows.append(" | ".join(cell.strip() for cell in row if cell is not None))
    return "\n".join(rows)


def extract_json(data: bytes) -> str:
    obj = json.loads(extract_txt(data))
    return json.dumps(obj, indent=2, ensure_ascii=False)


def extract_html(data: bytes) -> str:
    raw = extract_txt(data)
    raw = re.sub(r"(?is)<script.*?>.*?</script>", " ", raw)
    raw = re.sub(r"(?is)<style.*?>.*?</style>", " ", raw)
    raw = re.sub(r"(?is)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?is)</p>", "\n", raw)
    raw = re.sub(r"(?is)<[^>]+>", " ", raw)
    return html.unescape(re.sub(r"[ \t]+", " ", raw))


def extract_rtf(data: bytes) -> str:
    raw = extract_txt(data)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", raw).strip()


def extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_doc(data: bytes) -> str:
    """Old .doc — Word XML streams / utf-8 leftovers se best-effort text."""
    try:
        return extract_docx(data)
    except Exception:
        pass
    text = data.decode("latin-1", errors="ignore")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    text = re.sub(r"[^\x09\x0a\x0d\x20-\x7e\u00a0-\u024f]", " ", text)
    cleaned = re.sub(r"[ \t]{2,}", " ", text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    if len(cleaned.strip()) < 40:
        raise ValueError("Could not read .doc. Please save as .docx and upload again.")
    return cleaned


def extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def extract_xls(data: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    parts = []
    for sheet in book.sheets():
        parts.append(f"--- Sheet: {sheet.name} ---")
        for r in range(sheet.nrows):
            cells = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text)
    return "\n".join(parts)


def extract_image(data: bytes) -> str:
    """Gemini vision se image/scan se text."""
    from PIL import Image
    import google.generativeai as genai

    if not GEMINI_API_KEY:
        raise ValueError("Image OCR needs GEMINI_API_KEY")
    genai.configure(api_key=GEMINI_API_KEY)
    img = Image.open(io.BytesIO(data))
    model = genai.GenerativeModel("gemini-2.0-flash")
    response = model.generate_content(
        [
            "Extract ALL readable text from this image exactly. "
            "Keep names, numbers, tables as text. Return only the extracted text.",
            img,
        ]
    )
    text = response.text if hasattr(response, "text") else str(response)
    if not text or not text.strip():
        raise ValueError("No text found in image")
    return text


_HANDLERS: Dict[str, Callable[[bytes], str]] = {
    ".pdf": extract_pdf,
    ".txt": extract_txt,
    ".md": extract_txt,
    ".csv": extract_csv,
    ".json": extract_json,
    ".html": extract_html,
    ".htm": extract_html,
    ".rtf": extract_rtf,
    ".docx": extract_docx,
    ".doc": extract_doc,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".xls": extract_xls,
    ".pptx": extract_pptx,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
    ".webp": extract_image,
    ".gif": extract_image,
    ".bmp": extract_image,
    ".tiff": extract_image,
    ".tif": extract_image,
}


def extract_text(filename: str, data: bytes) -> str:
    ext = ext_of(filename)
    handler = _HANDLERS.get(ext)
    if not handler:
        raise ValueError(
            f"Unsupported file type: {ext or 'unknown'}. "
            f"Allowed: {', '.join(EXTRACTABLE_EXTENSIONS)}"
        )
    logger.info(f"📄 Extracting text from {filename} ({ext})")
    text = handler(data)
    text = (text or "").strip()
    if not text:
        raise ValueError(f"No text could be extracted from {filename}")
    logger.info(f"✅ Extracted {len(text)} chars from {filename}")
    return text
